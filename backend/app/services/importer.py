"""Turns uploaded files into page drafts (canvas content + extracted text).

PDF / photos   -> one page per PDF page; the page image is the background.
                  Typed PDF pages keep their extracted text; scanned or
                  handwritten pages are read later by the vision model.
PPTX           -> converted to PDF with LibreOffice when installed (slides look
                  exactly like the original); otherwise one text-only page per
                  slide (title + bullets as text boxes, speaker notes kept).
DOCX / TXT / MD -> text laid out on portrait pages as text boxes.
"""
import io
import os
import shutil
import subprocess
import tempfile
import textwrap
import uuid
from dataclasses import dataclass, field
from pathlib import Path

PAGE_W = 1200
MARGIN = 60
TEXT_MIN_CHARS = 40          # fewer extracted chars -> treat the page as an image
DOC_PAGE_H = 1600
IMG_EXTS = {".png", ".jpg", ".jpeg", ".bmp", ".gif", ".tif", ".tiff"}
FONT = {"h1": 40, "h2": 32, "body": 26, "bullet": 26}


class ImportProblem(Exception):
    pass


@dataclass
class PageDraft:
    title: str
    height: int
    objects: list = field(default_factory=list)
    background: bytes | None = None   # JPEG bytes of the page image
    source_text: str = ""             # text NOT already present as text boxes
    vision: bool = False              # needs the vision model to read it


# ------------------------------------------------------------------ helpers
def _fitz():
    try:
        import pymupdf as fitz
    except ImportError:
        try:
            import fitz  # older PyMuPDF
        except ImportError:
            raise ImportProblem("PDF/photo import needs PyMuPDF. Run: pip install pymupdf")
    return fitz


def _text_obj(x: float, y: float, text: str, size: int) -> dict:
    return {
        "id": uuid.uuid4().hex[:8],
        "type": "text",
        "x": x,
        "y": round(y),
        "text": text,
        "fontSize": size,
        "color": "#111111",
    }


def layout_blocks(blocks, paginate_height: int | None = None, min_height: int = 675):
    """blocks: [(kind, text, level)] with kind in h1|h2|body|bullet.
    Returns [(objects, text, height)]. Without paginate_height everything goes
    on one page that grows; with it, text flows across pages."""

    def new_page():
        return {"objects": [], "text": [], "y": float(MARGIN)}

    pages = [new_page()]
    for kind, text, level in blocks:
        text = (text or "").strip()
        if not text:
            continue
        fs = FONT[kind]
        lh = fs * 1.35
        indent = 0 if kind in ("h1", "h2", "body") else 28 + 28 * level
        x = MARGIN + indent
        chars = max(20, int((PAGE_W - MARGIN - x) / (fs * 0.52)))
        prefix = "\u2022 " if kind == "bullet" else ""
        lines = textwrap.wrap(
            prefix + " ".join(text.split()), chars, subsequent_indent="  " if prefix else ""
        ) or [""]
        if pages[-1]["objects"]:
            pages[-1]["y"] += 22 if kind in ("h1", "h2") else 10
        while lines:
            page = pages[-1]
            if paginate_height:
                room = int((paginate_height - MARGIN - page["y"]) // lh)
                if room < 1 and page["objects"]:
                    pages.append(new_page())
                    continue
                take = max(1, room)
            else:
                take = len(lines)
            chunk, lines = lines[:take], lines[take:]
            page["objects"].append(_text_obj(x, page["y"], "\n".join(chunk), fs))
            page["text"].append("\n".join(chunk))
            page["y"] += len(chunk) * lh
            if lines:
                pages.append(new_page())

    out = []
    for p in pages:
        if not p["objects"] and len(pages) > 1:
            continue
        height = paginate_height or max(min_height, int(p["y"] + MARGIN))
        out.append((p["objects"], "\n".join(p["text"]), height))
    return out


# ---------------------------------------------------------------------- PDF
def _pdf(data: bytes, stem: str, limit: int):
    fitz = _fitz()
    try:
        doc = fitz.open(stream=data, filetype="pdf")
    except Exception:
        raise ImportProblem("could not open this PDF (damaged?)")
    if doc.needs_pass:
        raise ImportProblem("this PDF is password protected")

    drafts: list[PageDraft] = []
    warnings: list[str] = []
    total = doc.page_count
    for i in range(min(total, limit)):
        page = doc[i]
        zoom = PAGE_W / max(page.rect.width, 1)
        pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
        text = page.get_text().strip()
        title = stem if total == 1 else f"{stem} - p{i + 1}"
        drafts.append(
            PageDraft(
                title=title[:200],
                height=max(300, min(pix.height, 4000)),
                background=pix.tobytes("jpeg"),
                source_text=text if len(text) >= TEXT_MIN_CHARS else "",
                vision=len(text) < TEXT_MIN_CHARS,
            )
        )
    if total > limit:
        warnings.append(f"{stem}: only the first {limit} of {total} pages were imported (page limit).")
    doc.close()
    return drafts, warnings


def _image(data: bytes, stem: str, ext: str, limit: int):
    fitz = _fitz()
    try:
        img_doc = fitz.open(stream=data, filetype=ext.lstrip("."))
        pdf_bytes = img_doc.convert_to_pdf()
    except Exception:
        raise ImportProblem("could not read this image (try PNG or JPEG)")
    return _pdf(pdf_bytes, stem, limit)


# --------------------------------------------------------------------- PPTX
def find_soffice() -> str | None:
    for name in ("soffice", "libreoffice"):
        found = shutil.which(name)
        if found:
            return found
    for p in (
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    ):
        if os.path.exists(p):
            return p
    return None


def _pptx_to_pdf(data: bytes) -> bytes | None:
    soffice = find_soffice()
    if not soffice:
        return None
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "slides.pptx"
        src.write_bytes(data)
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir", tmp, str(src)],
                capture_output=True,
                timeout=180,
                check=False,
            )
        except (subprocess.TimeoutExpired, OSError):
            return None
        out = Path(tmp) / "slides.pdf"
        return out.read_bytes() if out.exists() else None


def _iter_shapes(shapes):
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(shape.shapes)
        else:
            yield shape


def _pptx_text(data: bytes, stem: str, limit: int):
    try:
        from pptx import Presentation
    except ImportError:
        raise ImportProblem("PPTX import needs python-pptx. Run: pip install python-pptx")
    try:
        prs = Presentation(io.BytesIO(data))
    except Exception:
        raise ImportProblem("could not open this PPTX (damaged?)")

    ratio = (prs.slide_height / prs.slide_width) if prs.slide_width else 0.5625
    min_h = int(PAGE_W * ratio)
    slides = list(prs.slides)
    drafts: list[PageDraft] = []
    for n, slide in enumerate(slides[:limit], 1):
        title_shape = slide.shapes.title
        title_text = title_shape.text_frame.text.strip() if title_shape is not None and title_shape.has_text_frame else ""
        blocks = []
        if title_text:
            blocks.append(("h1", title_text, 0))
        shapes = sorted(
            _iter_shapes(slide.shapes),
            key=lambda s: (getattr(s, "top", None) or 0, getattr(s, "left", None) or 0),
        )
        for shape in shapes:
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = "".join(r.text for r in para.runs).strip()
                    if t:
                        kind = "bullet" if shape.is_placeholder else "body"
                        blocks.append((kind, t, para.level or 0))
            elif getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        blocks.append(("body", " | ".join(cells), 0))
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        pages = layout_blocks(blocks, paginate_height=None, min_height=min_h) or [([], "", min_h)]
        objs, _text, height = pages[0]
        drafts.append(
            PageDraft(
                title=(title_text[:60] or f"{stem} - slide {n}"),
                height=height,
                objects=objs,
                source_text=f"Speaker notes: {notes}" if notes else "",
            )
        )
    warnings = []
    if len(slides) > limit:
        warnings.append(f"{stem}: only the first {limit} of {len(slides)} slides were imported (page limit).")
    warnings.append(
        f"{stem}: imported as text (slide visuals/images are not included). "
        "For exact-looking slides, use PowerPoint > Save as PDF and import the PDF."
    )
    return drafts, warnings


def _pptx(data: bytes, stem: str, limit: int):
    pdf = _pptx_to_pdf(data)
    if pdf:
        try:
            return _pdf(pdf, stem, limit)
        except ImportProblem:
            pass
    return _pptx_text(data, stem, limit)


# --------------------------------------------------------------- DOCX / text
def _paged_text(blocks, stem: str, limit: int):
    pages = layout_blocks(blocks, paginate_height=DOC_PAGE_H)
    if not pages:
        raise ImportProblem("no text found in this file")
    drafts, warnings = [], []
    for i, (objs, _text, height) in enumerate(pages[:limit], 1):
        title = stem if len(pages) == 1 else f"{stem} ({i})"
        drafts.append(PageDraft(title=title[:200], height=height, objects=objs))
    if len(pages) > limit:
        warnings.append(f"{stem}: only the first {limit} of {len(pages)} pages were imported (page limit).")
    return drafts, warnings


def _docx(data: bytes, stem: str, limit: int):
    try:
        import docx
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError:
        raise ImportProblem("DOCX import needs python-docx. Run: pip install python-docx")
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception:
        raise ImportProblem("could not open this DOCX (damaged?)")

    blocks = []
    for el in document.element.body.iterchildren():
        tag = el.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(el, document)
            text = p.text.strip()
            if not text:
                continue
            style = (p.style.name if p.style is not None else "") or ""
            if style == "Title" or style == "Heading 1":
                blocks.append(("h1", text, 0))
            elif style.startswith("Heading"):
                blocks.append(("h2", text, 0))
            elif "List" in style:
                blocks.append(("bullet", text, 0))
            else:
                blocks.append(("body", text, 0))
        elif tag == "tbl":
            for row in Table(el, document).rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    blocks.append(("body", " | ".join(cells), 0))
    return _paged_text(blocks, stem, limit)


def _plain(data: bytes, stem: str, limit: int):
    blocks = []
    for raw in data.decode("utf-8", errors="replace").splitlines():
        line = raw.strip()
        if not line:
            continue
        if line.startswith("## "):
            blocks.append(("h2", line[3:], 0))
        elif line.startswith("#"):
            blocks.append(("h1", line.lstrip("#").strip(), 0))
        elif line[:2] in ("- ", "* "):
            blocks.append(("bullet", line[2:], 0))
        else:
            blocks.append(("body", line, 0))
    return _paged_text(blocks, stem, limit)


# ----------------------------------------------------------------- dispatch
def import_file(filename: str, data: bytes, limit: int):
    """Returns (drafts, warnings). Raises ImportProblem for a file we can't use."""
    ext = Path(filename).suffix.lower()
    stem = Path(filename).stem.strip()[:50] or "Imported"
    if ext == ".pdf":
        return _pdf(data, stem, limit)
    if ext in IMG_EXTS:
        return _image(data, stem, ext, limit)
    if ext == ".pptx":
        return _pptx(data, stem, limit)
    if ext == ".docx":
        return _docx(data, stem, limit)
    if ext in (".txt", ".md"):
        return _plain(data, stem, limit)
    if ext in (".ppt", ".doc"):
        raise ImportProblem("old .ppt/.doc files are not supported. Save as .pptx/.docx or export to PDF.")
    raise ImportProblem(f"unsupported file type '{ext or 'unknown'}' (use PDF, PPTX, DOCX, TXT, MD or an image)")
